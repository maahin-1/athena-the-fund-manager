from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from athena.fundamentals import export, page
from athena.fundamentals.rows import parse_rows
from athena.fundamentals.store import FundamentalsStore

FIXTURES = Path(__file__).parent / "fixtures" / "xbrl"


def doc_for(name: str, symbol: str, price: float | None = 3200.0) -> dict:
    with FundamentalsStore(":memory:") as store:
        statements = parse_rows((FIXTURES / f"{name}.xml").read_text(encoding="utf-8"))
        store.save_statements(statements, "u", datetime(2024, 4, 12, tzinfo=timezone.utc))
        store.save_company(symbol, f"{symbol} Ltd", "INE", "IT", statements.layout, statements.shares)
        return page.build_page(store, symbol, "consolidated", price, "warehouse close")


def test_the_workbook_has_a_sheet_per_section_with_numbers_as_numbers(tmp_path):
    doc = doc_for("tcs_fy2024", "TCS")
    out = tmp_path / "tcs.xlsx"
    export.write_workbook(doc, out)
    book = load_workbook(out)
    assert book.sheetnames == ["Summary", "Quarters", "Profit & Loss", "Balance Sheet", "Cash Flow", "Ratios", "Shareholding"]
    pl = book["Profit & Loss"]
    assert pl["A1"].value == "Rs crore" and pl["B1"].value == doc["annual"]["ends"][0]
    sales = next(r for r in doc["annual"]["income"] if r["key"] == "sales")
    row = next(r for r in pl.iter_rows(values_only=True) if r[0] == "Sales")
    assert row[1] == pytest.approx(sales["values"][0]) and isinstance(row[1], (int, float))


def test_a_missing_figure_is_an_empty_cell(tmp_path):
    doc = doc_for("tcs_fy2024", "TCS")
    out = tmp_path / "tcs.xlsx"
    export.write_workbook(doc, out)
    pl = load_workbook(out)["Profit & Loss"]
    assert pl.max_column >= 3  # the year column and the TTM column
    ttm = next(r for r in pl.iter_rows(values_only=True) if r[0] == "Sales")[-1]
    assert ttm is None  # one annual filing holds no four quarters, so no TTM, and the cell is empty rather than 0


def test_the_summary_names_the_source_the_basis_and_the_formulas(tmp_path):
    doc = doc_for("tcs_fy2024", "TCS")
    doc["gaps"] = ["balance sheets start at FY2023; the earlier years are not held yet"]
    out = tmp_path / "tcs.xlsx"
    export.write_workbook(doc, out)
    text = " ".join(str(c) for r in load_workbook(out)["Summary"].iter_rows(values_only=True) for c in r if c)
    assert "TCS" in text and "consolidated" in text and "NSE" in text
    assert "not held yet" in text
    assert "price x shares" in text  # a key figure's formula is written beside it


def test_a_bank_workbook_opens(tmp_path):
    doc = doc_for("hdfcbank_fy2024", "HDFCBANK", None)
    out = tmp_path / "bank.xlsx"
    export.write_workbook(doc, out)
    names = [r[0] for r in load_workbook(out)["Profit & Loss"].iter_rows(values_only=True)]
    assert "Pre-provision Profit" in names


def test_a_page_with_an_error_is_refused(tmp_path):
    with pytest.raises(ValueError, match="no statements"):
        export.write_workbook({"symbol": "X", "basis": "consolidated", "error": "no statements stored"}, tmp_path / "x.xlsx")
