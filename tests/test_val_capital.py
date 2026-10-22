from pathlib import Path

import pytest

from athena.valuation import capital as c
from athena.valuation import ratings as r

GINZU = Path(".athena/damodaran/site/pc/fcffsimpleginzu.xlsx")

# the sample company in Damodaran's fcffsimpleginzu.xlsx ('Cost of capital worksheet'): values read from the workbook's saved results
SHARES_X_PRICE = 311888.2
BOOK_DEBT, INTEREST, MATURITY = 45063.0, 493.4, 3.0
RISKFREE, TAX, UNLEVERED, ERP = 0.0458, 0.25, 0.46942377677075875, 0.05580814643676354
A1_SPREAD = 0.0070030000000000005


def test_the_synthetic_rating_matches_the_workbook_example_for_a_smaller_firm():
    # EBIT 3060.9 over interest 493.4 is a cover of 6.2037: A2/A in the smaller-firm table, spread 0.007751, cost of debt 0.0458 + 0.007751 + 0.0051 = 0.058651
    got = r.synthetic_rating(3060.9, 493.4, "small")
    assert (got.rating, round(got.coverage, 4), got.spread) == ("A2/A", 6.2037, 0.007751)
    debt = c.synthetic_cost_of_debt(3060.9, 493.4, 0.0458, "small", country_default_spread=0.0051)
    assert debt.pre_tax == pytest.approx(0.058651, abs=1e-9)
    assert r.synthetic_rating(3060.9, 493.4, "large").rating == "A1/A+"  # the same cover is rated higher for a large manufacturer: 5.5 to 6.5


def test_a_loss_gets_the_lowest_band_no_interest_the_highest_and_a_band_edge_belongs_to_the_band_it_starts():
    assert r.synthetic_rating(-50.0, 10.0).rating == "D2/D"
    assert r.synthetic_rating(50.0, 0.0).rating == "Aaa/AAA"
    assert r.synthetic_rating(6.0, 1.0, "small").rating == "A2/A"  # exactly 6.0 starts the A2/A band
    assert r.synthetic_rating(5.999, 1.0, "small").rating == "A3/A-"
    assert r.synthetic_rating(1.5, 1.0, "financial").rating == "A2/A"  # a lender needs far less cover for the same rating
    with pytest.raises(ValueError):
        r.synthetic_rating(1.0, 1.0, "huge")  # type: ignore[arg-type]


def test_the_cost_of_capital_reproduces_the_workbook_worksheet_to_the_last_digit():
    kd = RISKFREE + A1_SPREAD  # the sample uses its actual rating, A1/A+
    debt = c.market_value_of_debt(BOOK_DEBT, INTEREST, MATURITY, kd)
    assert debt == pytest.approx(39953.64461186393, rel=1e-12)
    result = c.cost_of_capital(c.CapitalInputs(RISKFREE, ERP, UNLEVERED, SHARES_X_PRICE, debt, TAX, kd))
    assert result.levered_beta == pytest.approx(0.5145245310221007, rel=1e-12)
    assert result.cost_of_equity == pytest.approx(0.07451466037258848, rel=1e-12)
    assert result.after_tax_cost_of_debt == pytest.approx(0.03960225, rel=1e-9)
    assert (result.equity_weight, result.debt_weight) == pytest.approx((0.8864443066573306, 0.11355569334266931), rel=1e-12)
    assert result.cost_of_capital == pytest.approx(0.0705501574064654, rel=1e-12)


def test_the_range_moves_the_cost_of_capital_up_and_down_and_keeps_the_base_where_it_was():
    base = c.CapitalInputs(0.05, 0.07, 0.9, 1000.0, 200.0, 0.25, 0.08)
    rng = c.capital_range(base, c.Bands(equity_risk_premium=0.01, riskfree=0.005, unlevered_beta=0.1))
    assert rng.low.cost_of_capital < rng.base.cost_of_capital < rng.high.cost_of_capital
    assert rng.base == c.cost_of_capital(base)
    assert rng.high.inputs.unlevered_beta == pytest.approx(1.0) and rng.low.inputs.equity_risk_premium == pytest.approx(0.06)


def test_rates_convert_between_currencies_by_the_inflation_gap_and_back():
    inr = c.convert_rate(0.083, inflation_to=0.04, inflation_from=0.02)
    assert inr == pytest.approx(1.083 * 1.04 / 1.02 - 1)
    assert c.convert_rate(inr, inflation_to=0.02, inflation_from=0.04) == pytest.approx(0.083)


def test_unusable_inputs_are_refused_not_guessed():
    with pytest.raises(ValueError):
        c.cost_of_capital(c.CapitalInputs(0.05, 0.07, 0.9, 0.0, 100.0, 0.25, 0.08))
    assert c.market_value_of_debt(100.0, 5.0, 3.0, 0.0) == 100.0  # no discount rate: face value


@pytest.mark.skipif(not GINZU.exists(), reason="the Damodaran mirror is not on this machine")
def test_the_tables_in_the_module_are_the_ones_in_the_workbook():
    import warnings

    import openpyxl

    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        sheet = openpyxl.load_workbook(GINZU, data_only=True)["Synthetic rating"]
    for kind, first_row in (("large", 22), ("small", 41)):
        rows = [(sheet.cell(row, 1).value, sheet.cell(row, 3).value, sheet.cell(row, 4).value) for row in range(first_row, first_row + 15)]
        mine = r.TABLES[kind]
        assert len(rows) == len(mine)
        for (lower, rating, spread), (m_lower, m_rating, m_spread) in zip(rows, mine):
            assert (m_rating, m_lower) == (rating, pytest.approx(lower)) and m_spread == pytest.approx(spread, abs=1e-6), (kind, rating)
