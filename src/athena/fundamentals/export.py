"""A company page as an Excel workbook: a summary, then one sheet per statement, periods across, figures as numbers.

    python -m athena.fundamentals.export TCS --out TCS.xlsx --basis consolidated --price 3200

The sheets hold what the page shows (rupee figures in crores). A figure the filings do not give is an empty cell, never zero."""
from __future__ import annotations

import argparse
import sys
from collections.abc import Sequence
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from athena.contracts import AthenaError
from athena.fundamentals.page import build_page
from athena.fundamentals.store import DEFAULT_PATH, open_read

BOLD = Font(bold=True)
SOURCE = "Source: NSE XBRL results filings. Rupee figures in crores (1 crore = 10,000,000)."


def _grid(sheet: Worksheet, corner: str, ends: Sequence[str], rows: Sequence[dict[str, Any]], *, ttm: bool = False, how: bool = False) -> None:
    sheet.append([corner, *ends, *(["TTM"] if ttm else []), *(["How it is worked out"] if how else [])])
    for cell in sheet[1]:
        cell.font = BOLD
    for r in rows:
        sheet.append([r["label"], *r["values"], *([r.get("ttm")] if ttm else []), *([r.get("how")] if how else [])])
    sheet.column_dimensions["A"].width = 34
    sheet.freeze_panes = "B2"


def _summary(sheet: Worksheet, doc: dict[str, Any]) -> None:
    sheet.append([f"{doc['name']} ({doc['symbol']})"])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([f"Basis: {doc['basis']}", f"Layout: {doc['layout']}", f"Latest period: {doc['as_of']}"])
    if doc.get("price") is not None:
        sheet.append([f"Price used: {doc['price']} ({doc.get('price_note') or 'supplied'})"])
    sheet.append([])
    sheet.append(["Key figures", "Value", "Unit", "How it is worked out"])
    for cell in sheet[sheet.max_row]:
        cell.font = BOLD
    for k in doc["key"]:
        sheet.append([k["label"], k["value"], k["unit"], k["how"]])
    for table in doc["growth"]:
        sheet.append([])
        sheet.append([table["title"], "%"])
        sheet[sheet.max_row][0].font = BOLD
        for r in table["rows"]:
            sheet.append([r["label"], r["value"]])
    if doc["gaps"]:
        sheet.append([])
        sheet.append(["Not held yet"])
        sheet[sheet.max_row][0].font = BOLD
        for gap in doc["gaps"]:
            sheet.append([gap])
    sheet.append([])
    sheet.append([SOURCE])
    sheet.append([f"Generated {date.today():%Y-%m-%d} by Athena."])
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["D"].width = 60
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top")


def write_workbook(doc: dict[str, Any], path: str | Path) -> None:
    if doc.get("error"):
        raise ValueError(f"cannot export {doc.get('symbol')}: {doc['error']}")
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    _summary(summary, doc)
    quarters, annual = doc["quarters"], doc["annual"]
    _grid(book.create_sheet("Quarters"), "Rs crore", quarters["ends"], quarters["rows"])
    _grid(book.create_sheet("Profit & Loss"), "Rs crore", annual["ends"], annual["income"], ttm=True)
    _grid(book.create_sheet("Balance Sheet"), "Rs crore", annual["ends"], annual["balance"])
    _grid(book.create_sheet("Cash Flow"), "Rs crore", annual["ends"], annual["cashflow"])
    _grid(book.create_sheet("Ratios"), "Ratio", doc["ratios"]["ends"], [*doc["ratios"]["rows"], *annual["banking"]], how=True)
    _grid(book.create_sheet("Shareholding"), "Percent of shares", doc["shareholding"]["ends"], doc["shareholding"]["rows"])
    book.save(str(path))


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("symbol")
    parser.add_argument("--out", required=True)
    parser.add_argument("--basis", default="consolidated", choices=("consolidated", "standalone"))
    parser.add_argument("--price", type=float, default=None)
    parser.add_argument("--price-note", default="")
    parser.add_argument("--store", default=str(DEFAULT_PATH))
    args = parser.parse_args(argv)
    try:
        with open_read(args.store) as store:
            doc = build_page(store, args.symbol.upper(), args.basis, args.price, args.price_note)
        write_workbook(doc, args.out)
    except (AthenaError, ValueError) as exc:
        print(str(exc), file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
